#!/usr/bin/env python3
"""Convert the Vita demo utterance workbook into a reusable intent bank.

The workbook (``Dataset - Vita Demo - <date>.xlsx``, sheet ``in``) is a balanced
stimulus grid: every subintent carries exactly one utterance per
``vehicle_state`` x ``ASSISTANT_MODE`` cell. This script keeps that grid intact
and reshapes it into JSON that both pipeline steps consume:

  step 1 (generation) — personas rephrase each seed in their own voice
  step 2 (scoring)    — the resulting utterances are sent to the Vita API

Usage:
    uv run python scripts/build_vita_intent_bank.py <workbook.xlsx> -o <out.json>
"""

from __future__ import annotations

import argparse
import json
from collections import Counter
from pathlib import Path
from typing import Any

import openpyxl

SHEET = "in"
COLUMNS = (
    "intent_code",
    "subintent_code",
    "subintent_name",
    "input",
    "vehicle_state",
    "ASSISTANT_MODE",
    "source",
    "note",
)


def read_rows(path: Path) -> list[dict[str, str]]:
    wb = openpyxl.load_workbook(path, data_only=True)
    if SHEET not in wb.sheetnames:
        raise SystemExit(f"sheet {SHEET!r} not found in {path} (has {wb.sheetnames})")
    grid = list(wb[SHEET].iter_rows(values_only=True))
    if not grid:
        raise SystemExit(f"{path} sheet {SHEET!r} is empty")

    header = [str(c).strip() if c is not None else "" for c in grid[0]]
    missing = [c for c in COLUMNS if c not in header]
    if missing:
        raise SystemExit(f"workbook is missing column(s): {', '.join(missing)}")
    index = {name: header.index(name) for name in COLUMNS}

    rows: list[dict[str, str]] = []
    for raw in grid[1:]:
        cell = {
            name: ("" if raw[pos] is None else str(raw[pos]).strip())
            for name, pos in index.items()
        }
        if not cell["input"]:
            continue  # trailing/blank spreadsheet rows
        rows.append(cell)
    return rows


def build_bank(rows: list[dict[str, str]]) -> dict[str, Any]:
    """Group flat rows into intent -> subintent -> seed utterances."""
    intents: dict[str, dict[str, Any]] = {}
    for row in rows:
        intent = intents.setdefault(
            row["intent_code"],
            {"intent_code": row["intent_code"], "subintents": {}},
        )
        subintent = intent["subintents"].setdefault(
            row["subintent_code"],
            {
                "subintent_code": row["subintent_code"],
                "subintent_name": row["subintent_name"],
                "seeds": [],
            },
        )
        subintent["seeds"].append(
            {
                "input": row["input"],
                "vehicle_state": row["vehicle_state"],
                "assistant_mode": row["ASSISTANT_MODE"],
            }
        )

    # Freeze dict ordering into lists so downstream batching is deterministic.
    return {
        "schemaVersion": "1.0",
        "artifactType": "vita.intent_bank",
        "intents": [
            {
                "intent_code": intent["intent_code"],
                "subintents": list(intent["subintents"].values()),
            }
            for intent in intents.values()
        ],
    }


def summarize(rows: list[dict[str, str]], bank: dict[str, Any]) -> str:
    states = Counter(r["vehicle_state"] for r in rows)
    modes = Counter(r["ASSISTANT_MODE"] for r in rows)
    lines = [
        f"seeds            : {len(rows)}",
        f"unique utterances: {len({r['input'] for r in rows})}",
        f"intents          : {len(bank['intents'])}",
        f"subintents       : {sum(len(i['subintents']) for i in bank['intents'])}",
        f"vehicle_state    : {dict(states)}",
        f"assistant_mode   : {dict(modes)}",
        "",
        "per intent:",
    ]
    for intent in bank["intents"]:
        count = sum(len(s["seeds"]) for s in intent["subintents"])
        lines.append(
            f"  {intent['intent_code']:28s} {count:3d} seeds"
            f"  {len(intent['subintents'])} subintents"
        )
    return "\n".join(lines)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path, help="path to the .xlsx workbook")
    parser.add_argument(
        "-o",
        "--output",
        type=Path,
        required=True,
        help="where to write the intent bank JSON",
    )
    args = parser.parse_args()

    rows = read_rows(args.workbook)
    bank = build_bank(rows)

    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        json.dumps(bank, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )

    print(summarize(rows, bank))
    print(f"\nwrote {args.output}")


if __name__ == "__main__":
    main()
